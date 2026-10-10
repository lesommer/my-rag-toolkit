from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from myrag.config import SUPPORTED_EXTENSIONS


@dataclass
class ParsedBlock:
    text: str
    section: str = ""
    page: int = -1
    kind: str = "text"


@dataclass
class ParsedDocument:
    doc_id: str
    source_path: str
    title: str
    fmt: str
    blocks: list[ParsedBlock] = field(default_factory=list)
    n_pages: int = 0
    scanned_pages: list[int] = field(default_factory=list)
    empty_pages: list[int] = field(default_factory=list)

    @property
    def is_scan(self) -> bool:
        return bool(self.n_pages) and len(self.scanned_pages) * 2 > self.n_pages


_PAGE_NUMBER_RE = re.compile(r"^(?:page\s*)?\d{1,4}(?:\s*/\s*\d{1,4})?$", re.IGNORECASE)
_HYPHEN_EOL_RE = re.compile(r"(\w)-\n(?=[a-z])")


def _norm_line(line: str) -> str:
    return re.sub(r"\d+", "#", line.lower()).strip()


def _clean_pdf_page(text: str) -> str:
    text = text.replace("\u00ad", "")
    lines = []
    for raw in text.split("\n"):
        line = re.sub(r"[ \t]+", " ", raw).strip()
        if line:
            lines.append(line)
    joined = "\n".join(lines)
    joined = _HYPHEN_EOL_RE.sub(r"\1", joined)
    joined = re.sub(r"\n{2,}", "\n\n", joined)
    return joined.strip()


def _recurring_lines(pages: list[str], n_pages: int) -> set[str]:
    if n_pages < 3:
        return set()
    threshold = max(3, round(n_pages * 0.6))
    counts: dict[str, int] = {}
    raws: dict[str, set[str]] = {}
    for page in pages:
        lines = [l.strip() for l in page.split("\n") if l.strip()]
        candidates = []
        if lines:
            candidates.append(lines[0])
        if len(lines) > 1:
            candidates.append(lines[-1])
        for cand in candidates:
            if len(cand) > 90:
                continue
            norm = _norm_line(cand)
            counts[norm] = counts.get(norm, 0) + 1
            raws.setdefault(norm, set()).add(cand)
    recurring: set[str] = set()
    for norm, count in counts.items():
        if count >= threshold:
            recurring |= raws[norm]
    return recurring


def _table_markdown(rows: list) -> str | None:
    clean: list[list[str]] = []
    for row in rows:
        cells = [re.sub(r"\s+", " ", c or "").strip() for c in row]
        if any(cells):
            clean.append(cells)
    if len(clean) < 2:
        return None
    width = max(len(r) for r in clean)
    lines = ["| " + " | ".join(clean[0] + [""] * (width - len(clean[0]))) + " |"]
    lines.append("|" + "---|" * width)
    for row in clean[1:]:
        lines.append("| " + " | ".join(row + [""] * (width - len(row))) + " |")
    return "\n".join(lines)


def _rect_inside(outer: tuple, inner: tuple, tol: float = 2.0) -> bool:
    return (
        inner[0] >= outer[0] - tol
        and inner[1] >= outer[1] - tol
        and inner[2] <= outer[2] + tol
        and inner[3] <= outer[3] + tol
    )


def _parse_pdf(path: Path) -> ParsedDocument:
    import pymupdf

    doc = pymupdf.open(path)
    meta = doc.metadata or {}
    title = (meta.get("title") or "").strip() or path.stem
    pages: list[str] = []
    extras: list[list[ParsedBlock]] = []
    url_entries: list[tuple[int, str, str]] = []
    scanned: list[int] = []
    empty: list[int] = []
    for i, page in enumerate(doc):
        raw = page.get_text("text") or ""
        if len(raw.strip()) < 20:
            if page.get_images(full=True):
                scanned.append(i + 1)
            else:
                empty.append(i + 1)
            pages.append("")
            extras.append([])
            continue
        table_rects: list[tuple] = []
        page_extras: list[ParsedBlock] = []
        try:
            for tab in page.find_tables().tables:
                try:
                    md = _table_markdown(tab.extract())
                except Exception:
                    md = None
                if md:
                    table_rects.append(tuple(tab.bbox))
                    page_extras.append(ParsedBlock(text=md, page=i + 1))
        except Exception:
            table_rects = []
            page_extras = []
        lines: list[str] = []
        try:
            for b in page.get_text("blocks"):
                if len(b) < 5 or not str(b[4]).strip():
                    continue
                rect = (b[0], b[1], b[2], b[3])
                if any(_rect_inside(tr, rect) for tr in table_rects):
                    continue
                lines.append(str(b[4]))
        except Exception:
            lines = [raw]
        try:
            for link in page.get_links():
                uri = str(link.get("uri") or "")
                if not uri.startswith(("http://", "https://")):
                    continue
                try:
                    anchor = re.sub(r"\s+", " ", page.get_text("text", clip=link["from"])).strip()
                except Exception:
                    anchor = ""
                url_entries.append((i + 1, uri, anchor[:80]))
        except Exception:
            pass
        pages.append(_clean_pdf_page("\n".join(lines)))
        extras.append(page_extras)
    n_pages = len(pages)
    doc.close()

    recurring = _recurring_lines(pages, n_pages)
    blocks: list[ParsedBlock] = []
    for i, page_text in enumerate(pages):
        page_blocks: list[ParsedBlock] = []
        if page_text:
            kept = []
            for line in page_text.split("\n"):
                stripped = line.strip()
                if not stripped:
                    continue
                if _PAGE_NUMBER_RE.match(stripped):
                    continue
                if stripped in recurring:
                    continue
                kept.append(stripped)
            body = "\n".join(kept).strip()
            if body:
                page_blocks.append(ParsedBlock(text=body, page=i + 1))
        page_blocks.extend(extras[i])
        blocks.extend(page_blocks)

    if url_entries:
        seen_urls: set[str] = set()
        inv_lines = ["URLs in this document:"]
        for page_no, uri, anchor in url_entries:
            if uri in seen_urls:
                continue
            seen_urls.add(uri)
            line = f"- p.{page_no}: {uri}"
            if anchor:
                line += f" [{anchor}]"
            inv_lines.append(line)
        blocks.append(ParsedBlock(text="\n".join(inv_lines), kind="urls"))

    return ParsedDocument(
        doc_id=path.name,
        source_path=str(path),
        title=title,
        fmt="pdf",
        blocks=blocks,
        n_pages=n_pages,
        scanned_pages=scanned,
        empty_pages=empty,
    )


_HEADING_STYLE_RE = re.compile(r"^(?:heading|titre)\s*(\d)", re.IGNORECASE)


def _parse_docx(path: Path) -> ParsedDocument:
    import docx as docxlib

    d = docxlib.Document(str(path))
    blocks: list[ParsedBlock] = []
    section = ""
    for para in d.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = ""
        if para.style is not None:
            style = para.style.name or ""
        if _HEADING_STYLE_RE.match(style.strip().lower()):
            section = text
        blocks.append(ParsedBlock(text=text, section=section))
    for table in d.tables:
        rows = []
        for row in table.rows:
            cells = [re.sub(r"\s+", " ", cell.text).strip() for cell in row.cells]
            rows.append("| " + " | ".join(cells) + " |")
        if rows:
            md = "\n".join(rows)
            if len(rows) >= 2:
                width = max(2, len(rows[0].split("|")) - 2)
                md = rows[0] + "\n|" + "---|" * width + "\n" + "\n".join(rows[1:])
            blocks.append(ParsedBlock(text=md, section=section))
    title = path.stem
    core = d.core_properties
    if core is not None and core.title:
        title = core.title
    return ParsedDocument(
        doc_id=path.name,
        source_path=str(path),
        title=title,
        fmt="docx",
        blocks=blocks,
    )


_MD_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
_MD_TABLE_RE = re.compile(r"^\s*\|.*\|\s*$")
_MD_IMAGE_RE = re.compile(r"^\s*!\[[^\]]*\]\([^)]*\)\s*$")


def _parse_markdown(path: Path, fmt: str) -> ParsedDocument:
    lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
    blocks: list[ParsedBlock] = []
    stack: list[tuple[int, str]] = []
    section = ""
    buf: list[str] = []
    table_buf: list[str] = []

    def flush_buf() -> None:
        if buf:
            text = "\n".join(buf).strip()
            if text:
                blocks.append(ParsedBlock(text=text, section=section))
            buf.clear()

    def flush_table() -> None:
        if table_buf:
            blocks.append(ParsedBlock(text="\n".join(table_buf), section=section))
            table_buf.clear()

    for line in lines:
        stripped = line.strip()
        if _MD_TABLE_RE.match(stripped):
            flush_buf()
            table_buf.append(stripped)
            continue
        flush_table()
        if not stripped:
            flush_buf()
            continue
        if _MD_IMAGE_RE.match(stripped):
            continue
        m = _MD_HEADING_RE.match(stripped)
        if m:
            flush_buf()
            level = len(m.group(1))
            while stack and stack[-1][0] >= level:
                stack.pop()
            stack.append((level, m.group(2)))
            section = " > ".join(t for _, t in stack)
            blocks.append(ParsedBlock(text=m.group(2), section=section))
            continue
        buf.append(stripped)
    flush_table()
    flush_buf()

    return ParsedDocument(
        doc_id=path.name,
        source_path=str(path),
        title=path.stem,
        fmt=fmt,
        blocks=blocks,
    )


def _parse_text(path: Path) -> ParsedDocument:
    raw = path.read_text(encoding="utf-8", errors="replace")
    blocks = []
    for para in re.split(r"\n\s*\n", raw):
        text = re.sub(r"[ \t]+", " ", para).strip()
        if text:
            blocks.append(ParsedBlock(text=text))
    return ParsedDocument(
        doc_id=path.name,
        source_path=str(path),
        title=path.stem,
        fmt="txt",
        blocks=blocks,
    )


def parse_file(path: Path) -> ParsedDocument:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _parse_pdf(path)
    if suffix == ".docx":
        return _parse_docx(path)
    if suffix in {".md", ".markdown"}:
        return _parse_markdown(path, "md")
    if suffix == ".txt":
        return _parse_text(path)
    raise ValueError(f"unsupported format: {suffix} ({path.name})")
